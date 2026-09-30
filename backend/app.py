"""AVELA: локальный сервер, API и SQLite. Запуск: python app.py."""
from pathlib import Path
from functools import wraps
from datetime import datetime, timedelta, timezone
from collections import defaultdict, deque
import io, json, os, re, secrets, sqlite3, threading, time, uuid
from urllib.parse import urlparse
from flask import Flask, request, jsonify, session, send_from_directory, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import requests
from pricing import SERVICES, RATES, PACKAGES, MATERIALS, EXTRAS, calculate, number

ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / '.env')
DATA = Path(os.environ.get('AVELA_DATA_DIR', ROOT / 'data'))
DATA.mkdir(parents=True, exist_ok=True)
UPLOADS = DATA / 'uploads'
UPLOADS.mkdir(exist_ok=True)
MSK = timezone(timedelta(hours=3))
PORT = int(os.getenv('PORT', '8000'))
secret_file = DATA / 'session.key'
if not secret_file.exists():
    secret_file.write_text(secrets.token_hex(48), encoding='utf-8')
    try: secret_file.chmod(0o600)
    except OSError: pass
app = Flask(__name__, static_folder=None)
app.config.update(SECRET_KEY=secret_file.read_text().strip(), MAX_CONTENT_LENGTH=51*1024*1024, SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Strict', SESSION_COOKIE_SECURE=os.getenv('COOKIE_SECURE') == '1', PERMANENT_SESSION_LIFETIME=timedelta(hours=8))

def db():
    conn = sqlite3.connect(DATA / 'avela.sqlite3', timeout=20)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA foreign_keys=ON')
    return conn

def now(): return datetime.now(MSK).isoformat(timespec='seconds')

def init_db():
    with db() as c:
        c.executescript('''
        PRAGMA journal_mode=WAL;
        CREATE TABLE IF NOT EXISTS admin (id INTEGER PRIMARY KEY CHECK(id=1), username TEXT NOT NULL, password TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS leads (id INTEGER PRIMARY KEY, created TEXT NOT NULL, data TEXT NOT NULL, telegram TEXT NOT NULL DEFAULT 'pending');
        CREATE TABLE IF NOT EXISTS files (id TEXT PRIMARY KEY, lead_id INTEGER NOT NULL REFERENCES leads(id), name TEXT NOT NULL, stored TEXT NOT NULL, mime TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS referrals (code TEXT PRIMARY KEY, name TEXT NOT NULL, created TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS events (id INTEGER PRIMARY KEY, created TEXT NOT NULL, event TEXT NOT NULL, visitor TEXT NOT NULL, visit TEXT NOT NULL, source TEXT NOT NULL, ref TEXT NOT NULL, detail TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS event_date ON events(created);
        CREATE INDEX IF NOT EXISTS event_visitor ON events(visitor);
        CREATE TABLE IF NOT EXISTS outbox (lead_id INTEGER PRIMARY KEY REFERENCES leads(id), next_part INTEGER DEFAULT 0, attempts INTEGER DEFAULT 0, next_try REAL DEFAULT 0);
        ''')
        if not c.execute('SELECT 1 FROM admin').fetchone():
            password = secrets.token_urlsafe(15)
            c.execute('INSERT INTO admin VALUES(1,?,?)', ('admin', generate_password_hash(password)))
            (DATA / 'first-login.txt').write_text(f'Логин: admin\nПароль: {password}\nАдмин-панель: http://localhost:{PORT}/admin\n', encoding='utf-8')
            try: (DATA / 'first-login.txt').chmod(0o600)
            except OSError: pass
            print('\nСоздан администратор. Логин и случайный пароль: data/first-login.txt\n', flush=True)
init_db()

# Запросы с других сайтов не должны создавать заявки или менять админ-панель.
@app.before_request
def protect():
    if request.method == 'POST':
        origin = request.headers.get('Origin')
        if origin and origin != request.host_url.rstrip('/'):
            return jsonify(error='Запрос с другого сайта запрещён.'), 403
        if request.headers.get('Sec-Fetch-Site') == 'cross-site':
            return jsonify(error='Межсайтовый запрос запрещён.'), 403

@app.after_request
def headers(response):
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    if request.path.startswith('/api') or request.path.startswith('/admin'):
        response.headers['Cache-Control'] = 'no-store'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self' https://mc.yandex.ru https://mc.yandex.com; style-src 'self'; img-src 'self' data: https://mc.yandex.ru https://mc.yandex.com; connect-src 'self' https://mc.yandex.ru https://mc.yandex.com; frame-src https://yandex.ru https://mc.yandex.ru; object-src 'none'; base-uri 'self'; form-action 'self'; frame-ancestors 'none'"
    return response

@app.errorhandler(413)
def too_big(_): return jsonify(error='Слишком большой запрос: до 5 файлов по 10 МБ.'), 413
@app.errorhandler(404)
def missing(_): return jsonify(error='Страница или файл не найдены.'), 404
@app.errorhandler(500)
def server_error(_): return jsonify(error='Ошибка сервера. Попробуйте ещё раз.'), 500

buckets = defaultdict(deque)
rate_lock = threading.Lock()
def limited(group, count, window):
    key = (group, request.remote_addr)
    current = time.time()
    with rate_lock:
        values = buckets[key]
        while values and values[0] < current-window: values.popleft()
        if len(values) >= count: return True
        values.append(current)
    return False

def admin_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get('admin'): return jsonify(error='Войдите в админ-панель.'), 401
        if request.method == 'POST' and not secrets.compare_digest(request.headers.get('X-CSRF-Token',''), session.get('csrf','missing')):
            return jsonify(error='Обновите страницу и повторите действие.'), 403
        return fn(*args, **kwargs)
    return wrapped

@app.get('/')
def index(): return send_from_directory(ROOT / 'public', 'index.html')
@app.get('/admin')
def admin_page(): return send_from_directory(ROOT / 'public', 'admin.html')
@app.get('/privacy.html')
def privacy(): return send_from_directory(ROOT / 'public', 'privacy.html')
@app.get('/<any(css,js,images):folder>/<path:filename>')
def assets(folder, filename): return send_from_directory(ROOT / 'public' / folder, filename)

@app.get('/api/config')
def config():
    content = json.loads((ROOT / 'content.json').read_text(encoding='utf-8'))
    mid = os.getenv('YANDEX_METRIKA_ID','')
    return jsonify(**content, services=SERVICES, rates=RATES, packages=PACKAGES, materials=MATERIALS, extras=EXTRAS, metrika=int(mid) if mid.isdigit() else None)

def short(raw, key, length=200): return str(raw.get(key,'')).strip()[:length]
def attribution(raw):
    if not isinstance(raw, dict): raise ValueError('Неверные данные источника.')
    a = {k:short(raw,k,500 if k=='referrer' else 150) for k in ['ref','source','medium','campaign','term','content','referrer','visitor','visit']}
    with db() as c:
        ref = c.execute('SELECT name FROM referrals WHERE code=?',(a['ref'],)).fetchone()
    a['ref_name'] = ref['name'] if ref else ''
    if not ref: a['ref']=''
    if not a['source']:
        host = urlparse(a['referrer']).hostname
        a['source'] = host if host and host != request.host.split(':')[0] else 'Прямой заход'
    return a

def valid_file(file):
    ext = Path(file.filename or '').suffix.lower()
    head = file.stream.read(32)
    file.stream.seek(0, 2); size = file.stream.tell(); file.stream.seek(0)
    matches = {'.pdf': head.startswith(b'%PDF-'), '.jpg':head.startswith(b'\xff\xd8\xff'), '.jpeg':head.startswith(b'\xff\xd8\xff'), '.png':head.startswith(b'\x89PNG\r\n\x1a\n'), '.gif':head.startswith((b'GIF87a',b'GIF89a')), '.webp':head.startswith(b'RIFF') and head[8:12]==b'WEBP'}
    if not matches.get(ext) or size == 0 or size > 10*1024*1024:
        raise ValueError('Файл должен быть PDF, JPG, PNG, GIF или WEBP, размером до 10 МБ. Расширение должно соответствовать содержимому.')
    mime = {'.pdf':'application/pdf','.jpg':'image/jpeg','.jpeg':'image/jpeg','.png':'image/png','.gif':'image/gif','.webp':'image/webp'}[ext]
    name = (file.filename or 'file').replace('\\','/').split('/')[-1]
    name = re.sub(r'[\x00-\x1f\x7f]', '', name)[:150]
    return ext, mime, name

@app.post('/api/leads')
def add_lead():
    if limited('lead', 8, 600): return jsonify(error='Слишком много заявок. Повторите через 10 минут.'), 429
    saved=[]
    try:
        f=request.form
        if f.get('website'): raise ValueError('Не удалось проверить форму.')
        name=short(f,'name',100); phone=short(f,'phone',30); email=short(f,'email',254); region=short(f,'region',150)
        if not name or not region: raise ValueError('Заполните имя и регион.')
        if not 10 <= len(re.sub(r'\D','',phone)) <= 15: raise ValueError('Телефон должен содержать от 10 до 15 цифр.')
        if email and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',email): raise ValueError('Проверьте электронную почту.')
        services=list(dict.fromkeys(f.getlist('services')))
        if not services or any(s not in SERVICES for s in services): raise ValueError('Выберите услуги из списка.')
        if f.get('consent') != 'on': raise ValueError('Нужно согласие на обработку данных.')
        area=number(f['area'],1,3000,'Площадь') if f.get('area') else None
        calc=calculate(json.loads(f['calculator'])) if f.get('calculator') else None
        attr=attribution(json.loads(f.get('attribution','{}')))
        files=[file for file in request.files.getlist('files') if file.filename]
        if len(files)>5: raise ValueError('Можно прикрепить до 5 файлов.')
        validated=[(file,valid_file(file)) for file in files]
        payload={'name':name,'phone':phone,'email':email,'region':region,'services':services,'area':area,'comment':short(f,'comment',5000),'calculator':calc,'attribution':attr,'consent':True,'consent_version':'2026-09-29','consent_at':now()}
        with db() as c:
            lead_id=c.execute('INSERT INTO leads(created,data) VALUES(?,?)',(now(),json.dumps(payload,ensure_ascii=False))).lastrowid
            for file,(ext,mime,name) in validated:
                fid=uuid.uuid4().hex; stored=fid+ext; path=UPLOADS/stored
                file.save(path);saved.append(path)
                c.execute('INSERT INTO files VALUES(?,?,?,?,?)',(fid,lead_id,name,stored,mime))
            c.execute('INSERT INTO outbox(lead_id) VALUES(?)',(lead_id,))
        return jsonify(id=lead_id),201
    except (ValueError,TypeError,KeyError) as exc:
        for path in saved: path.unlink(missing_ok=True)
        return jsonify(error=str(exc)),400
    except Exception:
        for path in saved: path.unlink(missing_ok=True)
        raise

@app.post('/api/event')
def event():
    if limited('event',150,60): return jsonify(error='Лимит событий.'),429
    try:
        raw=request.get_json(silent=True) or {}
        if not isinstance(raw,dict): raise ValueError('Некорректное событие.')
        event=short(raw,'event',30)
        if event not in ['pageview','phone','email','cta','calculator','project','lead']: raise ValueError('Неизвестное событие.')
        a=attribution(raw)
        if not all(re.fullmatch(r'[a-zA-Z0-9-]{16,80}',a[x]) for x in ['visitor','visit']): raise ValueError('Нет идентификатора посещения.')
        with db() as c:
            c.execute('INSERT INTO events(created,event,visitor,visit,source,ref,detail) VALUES(?,?,?,?,?,?,?)',(now(),event,a['visitor'],a['visit'],a['source'],a['ref'],short(raw,'detail',200)))
        return jsonify(ok=True)
    except ValueError as exc: return jsonify(error=str(exc)),400

@app.post('/api/login')
def login():
    if limited('login',5,300): return jsonify(error='Слишком много попыток. Подождите 5 минут.'),429
    raw=request.get_json(silent=True) or {}
    if not isinstance(raw,dict): return jsonify(error='Неверный запрос.'),400
    with db() as c: admin=c.execute('SELECT * FROM admin WHERE id=1').fetchone()
    if short(raw,'username',100)!=admin['username'] or not check_password_hash(admin['password'],short(raw,'password',300)):
        return jsonify(error='Неверный логин или пароль.'),401
    session.clear();session['admin']=True;session['csrf']=secrets.token_hex(32);session.permanent=True
    return jsonify(csrf=session['csrf'])
@app.get('/api/admin/session')
@admin_required
def check_session(): return jsonify(csrf=session['csrf'])
@app.post('/api/admin/logout')
@admin_required
def logout(): session.clear();return jsonify(ok=True)

@app.get('/api/admin/leads')
@admin_required
def leads():
    with db() as c:
        rows=c.execute('SELECT * FROM leads ORDER BY id DESC').fetchall()
        allfiles=c.execute('SELECT id,lead_id,name FROM files').fetchall()
    return jsonify([{'id':r['id'],'created':r['created'],'telegram':r['telegram'],**json.loads(r['data']),'files':[dict(f) for f in allfiles if f['lead_id']==r['id']]} for r in rows])
@app.get('/api/admin/files/<fid>')
@admin_required
def download(fid):
    with db() as c: f=c.execute('SELECT * FROM files WHERE id=?',(fid,)).fetchone()
    if not f: return jsonify(error='Файл не найден.'),404
    return send_from_directory(UPLOADS,f['stored'],as_attachment=True,download_name=f['name'],mimetype=f['mime'])

@app.post('/api/admin/referrals')
@admin_required
def create_ref():
    raw=request.get_json(silent=True) or {}
    name=short(raw,'name',150) if isinstance(raw,dict) else ''
    if not name: return jsonify(error='Введите название источника.'),400
    with db() as c:
        while True:
            code=secrets.token_urlsafe(6)
            try: c.execute('INSERT INTO referrals VALUES(?,?,?)',(code,name,now()));break
            except sqlite3.IntegrityError: continue
    return jsonify(code=code,name=name,url=request.host_url+'?ref='+code),201

def bounds(period):
    n=datetime.now(MSK);today=n.replace(hour=0,minute=0,second=0,microsecond=0)
    if period=='yesterday': return today-timedelta(days=1),today
    starts={'today':today,'7':today-timedelta(days=6),'30':today-timedelta(days=29),'month':today.replace(day=1),'year':today.replace(month=1,day=1)}
    return starts.get(period,starts['30']),n+timedelta(seconds=1)

def week(date):
    d=datetime.fromisoformat(date).astimezone(MSK)
    return (d-timedelta(days=d.weekday())).date().isoformat()

def stats(events, leads):
    views=[e for e in events if e['event']=='pageview']
    visitors={e['visitor'] for e in views}; visits={e['visit'] for e in views}
    converted={l['attribution']['visitor'] for l in leads if l['attribution']['visitor'] in visitors}
    return {'visits':len(visits),'unique':len(visitors),'leads':len(leads),'conversion':round(100*len(converted)/len(visitors),2) if visitors else 0,'phone':sum(e['event']=='phone' for e in events),'email':sum(e['event']=='email' for e in events),'cta':sum(e['event']=='cta' for e in events),'pageviews':len(views)}

@app.get('/api/admin/analytics')
@admin_required
def analytics():
    start,end=bounds(request.args.get('period','30'))
    with db() as c:
        events=[dict(r) for r in c.execute('SELECT * FROM events WHERE created>=? AND created<?',(start.isoformat(),end.isoformat()))]
        leads=[{'id':r['id'],'created':r['created'],**json.loads(r['data'])} for r in c.execute('SELECT * FROM leads WHERE created>=? AND created<?',(start.isoformat(),end.isoformat()))]
        referrals=[dict(r) for r in c.execute('SELECT * FROM referrals ORDER BY created DESC')]
        allviews=[dict(r) for r in c.execute("SELECT created,visitor FROM events WHERE event='pageview' ORDER BY created")]
    sources=sorted({e['source'] for e in events}|{l['attribution']['source'] for l in leads})
    source_stats=[{'name':s,**stats([e for e in events if e['source']==s],[l for l in leads if l['attribution']['source']==s])} for s in sources]
    for r in referrals:
        ev=[e for e in events if e['ref']==r['code']];ls=[l for l in leads if l['attribution']['ref']==r['code']]
        r.update(stats(ev,ls));r['url']=request.host_url+'?ref='+r['code']
    weeks=sorted({week(e['created']) for e in events}|{week(l['created']) for l in leads})
    weekly=[{'week':w,**stats([e for e in events if week(e['created'])==w],[l for l in leads if week(l['created'])==w])} for w in weeks]
    # Когорты по неделе первого визита. Ячейки = вернувшиеся уникальные посетители.
    first={};active=defaultdict(set)
    for e in allviews:
        w=week(e['created']); first.setdefault(e['visitor'],w);active[e['visitor']].add(w)
    grouped=defaultdict(list)
    for v,w in first.items():
        if week(start.isoformat())<=w<=week((end-timedelta(seconds=1)).isoformat()):grouped[w].append(v)
    cohorts=[]
    current_week=week(now())
    for w,visitors in sorted(grouped.items()):
        values=[]
        for i in range(4):
            target=(datetime.fromisoformat(w)+timedelta(weeks=i)).date().isoformat()
            values.append(None if target>current_week else round(100*sum(target in active[v] for v in visitors)/len(visitors),1))
        cohorts.append({'week':w,'size':len(visitors),'retention':values})
    return jsonify(summary=stats(events,leads),sources=source_stats,referrals=referrals,weekly=weekly,cohorts=cohorts,period_start=start.isoformat(),period_end=end.isoformat(),telegram_configured=bool(os.getenv('TELEGRAM_BOT_TOKEN') and os.getenv('TELEGRAM_CHAT_ID')))

@app.get('/api/admin/export')
@admin_required
def export():
    wb=Workbook();ws=wb.active;ws.title='Заявки'
    titles=['ID','Дата (МСК)','Имя','Телефон','Email','Регион','Услуги','Площадь','Комментарий','Калькулятор — все параметры','Предварительная сумма','Файлы — имена и URL','Источник','Реферальное название','Реферальный код','UTM medium','UTM campaign','UTM term','UTM content','Referer','Visitor','Visit','Согласие','Дата согласия','Версия согласия','Telegram']
    ws.append(titles)
    with db() as c:
        for row in c.execute('SELECT * FROM leads ORDER BY id DESC'):
            d=json.loads(row['data']);a=d['attribution'];calc=d['calculator']
            files=c.execute('SELECT id,name FROM files WHERE lead_id=?',(row['id'],)).fetchall()
            values=[row['id'],row['created'],d['name'],d['phone'],d['email'],d['region'],', '.join(d['services']),d['area'],d['comment'],json.dumps(calc,ensure_ascii=False) if calc else '',calc['total'] if calc else '', '\n'.join(f['name']+' | '+request.host_url+'api/admin/files/'+f['id'] for f in files),a['source'],a['ref_name'],a['ref'],a['medium'],a['campaign'],a['term'],a['content'],a['referrer'],a['visitor'],a['visit'],'Да',d['consent_at'],d['consent_version'],row['telegram']]
            ws.append(values)
            # Не позволяем пользовательскому тексту стать формулой Excel.
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str): cell.data_type='s'
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='293B31')
    for column in ws.columns: ws.column_dimensions[column[0].column_letter].width=26
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    buf=io.BytesIO();wb.save(buf);buf.seek(0)
    return send_file(buf,as_attachment=True,download_name='AVELA_заявки.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# Доставка Telegram выполняется в фоне. Уже отправленные части отмечаются в SQLite.
def deliver(row):
    token=os.getenv('TELEGRAM_BOT_TOKEN');chat=os.getenv('TELEGRAM_CHAT_ID')
    if not token or not chat: return
    with db() as c:
        lead=c.execute('SELECT * FROM leads WHERE id=?',(row['lead_id'],)).fetchone()
        files=c.execute('SELECT * FROM files WHERE lead_id=?',(row['lead_id'],)).fetchall()
    d=json.loads(lead['data'])
    message=f"Новая заявка AVELA №{lead['id']}\n{lead['created']}\n"+json.dumps(d,ensure_ascii=False,indent=2)
    chunks=[message[i:i+2800] for i in range(0,len(message),2800)]
    parts=[('text',t) for t in chunks]+[('file',f) for f in files]
    for i in range(row['next_part'],len(parts)):
        kind,item=parts[i]
        try:
            if kind=='text': response=requests.post(f'https://api.telegram.org/bot{token}/sendMessage',json={'chat_id':chat,'text':item},timeout=30)
            else:
                with (UPLOADS/item['stored']).open('rb') as f:
                    response=requests.post(f'https://api.telegram.org/bot{token}/sendDocument',data={'chat_id':chat,'caption':f"Заявка AVELA №{lead['id']}"},files={'document':(item['name'],f,item['mime'])},timeout=60)
            result=response.json()
            if not response.ok or not result.get('ok'): raise RuntimeError('Telegram error')
            with db() as c: c.execute('UPDATE outbox SET next_part=? WHERE lead_id=?',(i+1,row['lead_id']))
        except Exception:
            attempts=row['attempts']+1
            with db() as c:
                c.execute('UPDATE outbox SET attempts=?,next_try=? WHERE lead_id=?',(attempts,time.time()+min(3600,30*2**min(attempts,7)),row['lead_id']))
                c.execute("UPDATE leads SET telegram='retry' WHERE id=?",(row['lead_id'],))
            return
    with db() as c:
        c.execute("UPDATE leads SET telegram='sent' WHERE id=?",(row['lead_id'],));c.execute('DELETE FROM outbox WHERE lead_id=?',(row['lead_id'],))

def worker():
    while True:
        try:
            if os.getenv('TELEGRAM_BOT_TOKEN') and os.getenv('TELEGRAM_CHAT_ID'):
                with db() as c: rows=c.execute('SELECT * FROM outbox WHERE next_try<=? ORDER BY lead_id LIMIT 10',(time.time(),)).fetchall()
                for row in rows: deliver(row)
        except Exception: app.logger.error('Ошибка очереди Telegram. Повторим позже.')
        time.sleep(5)

if __name__=='__main__':
    threading.Thread(target=worker,daemon=True).start()
    print(f'Сайт: http://localhost:{PORT}\nАдмин-панель: http://localhost:{PORT}/admin',flush=True)
    app.run(host='127.0.0.1',port=PORT,debug=False)
